"""Evidence-based, read-only scans for CSV and Excel files.

Findings are candidates, not confirmed exploits. The scanner never executes spreadsheet
formulas, changes source data, connects to databases, or sends raw rows to an LLM.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import pandas as pd

from gradientharbor.security.crypto import sign_bytes, verify_signature
from gradientharbor.security.prompt_guard import scan_untrusted_text

SENSITIVE_NAME = re.compile(r"(?i)(password|passwd|secret|api[_ -]?key|access[_ -]?token|private[_ -]?key|ssn|social[_ -]?security|credit[_ -]?card|card[_ -]?number|bank[_ -]?account|email|e-mail|phone|mobile|date[_ -]?of[_ -]?birth|dob)")
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
PHONE_RE = re.compile(r"(?<!\w)(?:\+?\d[\d ()-]{8,}\d)(?!\w)")
CARD_RE = re.compile(r"(?<!\d)(?:\d[ -]?){13,19}(?!\d)")
SECRET_RE = re.compile(r"(?i)(?:nvapi-[A-Za-z0-9_-]{20,}|\bAKIA[0-9A-Z]{16}\b|-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----|\b(?:api[_ -]?key|access[_ -]?token|password|secret[_ -]?key)\s*[:=]\s*['\"]?[^\s'\"]{12,})")
FORMULA_PREFIX = re.compile(r"^\s*[=+@]")

SEVERITY_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Info": 4}


def _finding(category: str, title: str, severity: str, confidence: str, evidence: str,
             remediation: str, *, rows: int | None = None, status: str = "Candidate — needs review",
             finding_id: str | None = None) -> dict[str, Any]:
    stable = f"{category}|{title}|{evidence}|{rows or 0}"
    return {
        "finding_id": finding_id or hashlib.sha256(stable.encode("utf-8")).hexdigest()[:12],
        "category": category, "title": title, "severity": severity, "confidence": confidence,
        "evidence": evidence, "affected_rows": rows, "status": status, "remediation": remediation,
    }


def load_tabular_file(file_bytes: bytes, filename: str) -> tuple[dict[str, pd.DataFrame], dict[str, Any]]:
    """Load CSV/XLSX without evaluating formulas. Returns frames and workbook metadata."""
    if not file_bytes:
        raise ValueError("The uploaded file is empty.")
    if len(file_bytes) > 50 * 1024 * 1024:
        raise ValueError("File exceeds the 50 MiB read-only scan limit. Split the dataset into smaller files.")
    name = filename.lower()
    metadata: dict[str, Any] = {"format": "csv" if name.endswith(".csv") else "excel", "sheets": [], "hidden_sheets": [], "formula_cells": 0, "external_links": 0}
    if name.endswith(".csv"):
        try:
            frame = pd.read_csv(BytesIO(file_bytes), low_memory=False)
        except Exception as exc:
            raise ValueError(f"CSV could not be parsed: {type(exc).__name__}: {exc}") from exc
        metadata["sheets"] = ["CSV"]
        return {"CSV": frame}, metadata
    if not name.endswith((".xlsx", ".xlsm")):
        raise ValueError("Supported formats are CSV, XLSX, and XLSM. Legacy XLS is not supported by this scanner.")
    try:
        # XLSX/XLSM are ZIP containers. Reject suspicious expansion ratios before parsing.
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            members = archive.infolist()
            total_uncompressed = sum(item.file_size for item in members)
            largest_ratio = max((item.file_size / max(item.compress_size, 1) for item in members), default=1.0)
            if total_uncompressed > 200 * 1024 * 1024 or largest_ratio > 200:
                raise ValueError("Workbook exceeds safe decompression limits and was not opened.")
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=False, data_only=False, keep_links=True)
        metadata["sheets"] = list(wb.sheetnames)
        metadata["hidden_sheets"] = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
        metadata["external_links"] = len(getattr(wb, "_external_links", []) or [])
        formula_count = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_count += 1
        metadata["formula_cells"] = formula_count
        wb.close()
        frames: dict[str, pd.DataFrame] = {}
        excel = pd.ExcelFile(BytesIO(file_bytes), engine="openpyxl")
        for sheet in excel.sheet_names:
            frames[sheet] = pd.read_excel(excel, sheet_name=sheet, dtype=object)
        excel.close()
        return frames, metadata
    except Exception as exc:
        raise ValueError(f"Excel workbook could not be parsed: {type(exc).__name__}: {exc}") from exc


def scan_dataframe(df: pd.DataFrame, sheet_name: str = "CSV", max_text_cells: int = 20000) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    if df.empty and len(df.columns) == 0:
        return [_finding("Data quality & poisoning", "Empty dataset", "Medium", "High",
                         f"Sheet '{sheet_name}' contains no rows or columns.", "Confirm the correct file and export range were uploaded.")]

    rows, cols = df.shape
    if rows == 0:
        findings.append(_finding("Data quality & poisoning", "No data rows", "Medium", "High",
                                 f"Sheet '{sheet_name}' has {cols} columns but no data rows.", "Confirm the source export completed successfully."))
    if cols == 0:
        findings.append(_finding("Data quality & poisoning", "No columns detected", "High", "High",
                                 f"Sheet '{sheet_name}' has no detected columns.", "Re-export the source file with a header row."))
        return findings

    dup_count = int(df.duplicated(keep="first").sum()) if rows else 0
    if dup_count:
        findings.append(_finding("Data quality & poisoning", "Duplicate rows detected", "Low" if dup_count / max(rows, 1) < .01 else "Medium", "High",
                                 f"{dup_count} row(s) duplicate an earlier row in sheet '{sheet_name}'. No row values are copied into this report.",
                                 "Review business keys and determine whether duplicates are valid repeated events or accidental copies.", rows=dup_count))

    missing = df.isna().sum()
    for col, count in missing.items():
        pct = float(count / max(rows, 1) * 100)
        if count and pct >= 20:
            findings.append(_finding("Data quality & poisoning", "High missing-value rate", "Medium" if pct < 60 else "High", "High",
                                     f"Column '{str(col)[:100]}' has {int(count)} missing value(s) ({pct:.1f}%).",
                                     "Check upstream ingestion, optional-field rules, and whether missingness is systematic.", rows=int(count)))

    for col in df.columns:
        series = df[col].dropna()
        if len(series) >= 10 and series.nunique(dropna=True) == 1:
            findings.append(_finding("Data quality & poisoning", "Constant-valued column", "Low", "High",
                                     f"Column '{str(col)[:100]}' has one distinct non-null value across {len(series)} value(s).",
                                     "Confirm the column is intended to be constant and not stuck due to a failed import."))
        # Numeric outliers are only candidates; they do not prove poisoning.
        numeric = pd.to_numeric(series, errors="coerce").dropna()
        if len(numeric) >= 12 and numeric.nunique() >= 2:
            q1, q3 = numeric.quantile(.25), numeric.quantile(.75)
            iqr = q3 - q1
            outlier_method = "3×IQR fences"
            if pd.notna(iqr) and iqr > 0:
                lower, upper = q1 - 3 * iqr, q3 + 3 * iqr
                outlier_mask = (numeric < lower) | (numeric > upper)
            else:
                # Handle common cases where most values are identical and IQR is zero.
                median = float(numeric.median())
                deviations = (numeric - median).abs()
                mad = float(deviations.median())
                if mad > 0:
                    modified_z = 0.6745 * (numeric - median) / mad
                    outlier_mask = modified_z.abs() > 3.5
                    outlier_method = "modified z-score > 3.5"
                else:
                    threshold = max(abs(median) * 10.0, 1e-9)
                    outlier_mask = deviations > threshold
                    outlier_method = "10× median-deviation fallback (zero IQR/MAD)"
            outliers = int(outlier_mask.sum())
            if outliers and outliers / len(numeric) <= .10:
                findings.append(_finding("Data quality & poisoning", "Extreme numeric outlier candidate", "Medium", "Medium",
                                         f"Column '{str(col)[:100]}' has {outliers} value(s) flagged by {outlier_method}. This is an anomaly signal, not proof of data poisoning.",
                                         "Validate these values against the source system, expected ranges, and change history.", rows=outliers))

    # Sensitive-column names and value patterns; do not expose actual cell contents.
    for col in df.columns:
        col_name = str(col)
        series = df[col].dropna().astype(str).head(10000)
        if SENSITIVE_NAME.search(col_name):
            findings.append(_finding("Database/data exposure", "Potentially sensitive column", "Medium", "Medium",
                                     f"Column name '{col_name[:100]}' suggests personal or secret data. Values are not copied into the report.",
                                     "Confirm the data classification, minimize access, and mask or tokenize values not needed for analysis."))
        secret_rows = [i for i, value in enumerate(series) if SECRET_RE.search(value)]
        if secret_rows:
            findings.append(_finding("Database/data exposure", "Credential-shaped value detected", "High", "High",
                                     f"Column '{col_name[:100]}' contains {len(secret_rows)} value(s) matching common credential patterns; sample values are intentionally redacted.",
                                     "Treat any real credential as exposed: revoke/rotate it, remove it from exports, and inspect access history.", rows=len(secret_rows)))
        email_rows = sum(1 for value in series if EMAIL_RE.search(value))
        if email_rows and not re.search(r"(?i)(email|e-mail)", col_name):
            findings.append(_finding("Database/data exposure", "Email-like values in a non-email column", "Low", "Medium",
                                     f"Column '{col_name[:100]}' has {email_rows} email-like value(s); values are not shown.",
                                     "Confirm whether this column contains personal data and apply appropriate minimization or masking.", rows=email_rows))
        card_rows = sum(1 for value in series if CARD_RE.search(value) and len(re.sub(r"\D", "", value)) in range(13, 20))
        if card_rows:
            findings.append(_finding("Database/data exposure", "Payment-card-like numeric values", "Medium", "Low",
                                     f"Column '{col_name[:100]}' has {card_rows} value(s) matching a card-number shape. This scan does not validate card numbers or PCI scope.",
                                     "Check with an approved payment-data scanner and avoid storing raw payment-card data.", rows=card_rows))

    # Spreadsheet formula injection and indirect prompt injection are scanned in cells.
    formula_rows = []
    injection_cells = []
    scanned = 0
    text_cols = [c for c in df.columns if df[c].dtype == object or pd.api.types.is_string_dtype(df[c])]
    for col in text_cols:
        for row_pos, value in enumerate(df[col].head(max_text_cells).tolist(), start=2):
            if not isinstance(value, str) or not value:
                continue
            scanned += 1
            if FORMULA_PREFIX.match(value):
                formula_rows.append((row_pos, str(col)))
            if len(injection_cells) < 100:
                guard = scan_untrusted_text(value[:3000])
                if guard.get("suspicious"):
                    injection_cells.append((row_pos, str(col), guard.get("severity", "medium"), guard.get("signals", [])))
    if formula_rows:
        findings.append(_finding("Database/data exposure", "Spreadsheet formula-injection candidate", "High", "Medium",
                                 f"{len(formula_rows)} text cell(s) begin with formula-like prefixes (=, +, or @) in sheet '{sheet_name}'. This does not execute any formula.",
                                 "If exporting to CSV or opening in spreadsheet software, neutralize untrusted cells (for example, prefix with an apostrophe) and validate downstream import behavior.", rows=len(formula_rows)))
    if injection_cells:
        unique_signals = sorted({sig for _, _, _, signals in injection_cells for sig in signals})
        high = sum(1 for _, _, severity, _ in injection_cells if severity == "high")
        findings.append(_finding("AI-agent & prompt injection", "Prompt-injection-like text in dataset", "High" if high else "Medium", "Medium",
                                 f"Detected {len(injection_cells)} suspicious text cell(s) among up to {scanned} scanned text cells in '{sheet_name}'. Signals: {', '.join(unique_signals[:8])}. Cell contents are not copied into the report.",
                                 "Treat these cells as untrusted data, keep them separate from system instructions, and require policy checks before retrieval or tool use.", rows=len(injection_cells)))
    return findings


def scan_workbook_metadata(metadata: dict[str, Any]) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    hidden = metadata.get("hidden_sheets") or []
    if hidden:
        findings.append(_finding("Database/data exposure", "Hidden worksheet(s) present", "Low", "High",
                                 f"Workbook contains {len(hidden)} hidden or very-hidden worksheet(s): {', '.join(str(x)[:80] for x in hidden[:10])}.",
                                 "Review hidden worksheets before sharing or importing the workbook; hidden does not necessarily mean malicious."))
    if metadata.get("external_links", 0):
        findings.append(_finding("Database/data exposure", "External workbook links present", "Medium", "High",
                                 f"Workbook metadata reports {metadata['external_links']} external link(s). Links were not opened or followed.",
                                 "Review and remove unnecessary external links; do not refresh links from untrusted workbooks."))
    if metadata.get("formula_cells", 0):
        findings.append(_finding("Database/data exposure", "Workbook contains formulas", "Info", "High",
                                 f"Detected {metadata['formula_cells']} formula cell(s). The scanner reads formulas but never evaluates them.",
                                 "Review formulas in untrusted workbooks and inspect external references or volatile functions if relevant."))
    return findings


def file_sha256(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def _manifest_payload(filename: str, digest: str, created_at: str) -> bytes:
    return json.dumps({"schema": "gradientharbor.integrity-baseline.v1", "filename": filename,
                       "sha256": digest, "created_at": created_at}, sort_keys=True, separators=(",", ":")).encode("utf-8")


def create_signed_baseline(filename: str, file_bytes: bytes) -> dict[str, Any]:
    """Sign a snapshot for future comparisons; this does not prove pre-upload authenticity."""
    digest = file_sha256(file_bytes)
    created_at = datetime.now(timezone.utc).isoformat()
    payload = _manifest_payload(filename, digest, created_at)
    return {"schema": "gradientharbor.integrity-baseline.v1", "filename": filename,
            "sha256": digest, "created_at": created_at, "signature_ed25519": sign_bytes(payload),
            "note": "Trust this baseline only if it was created from a trusted file and stored separately from the file being checked."}


def verify_signed_baseline(manifest: dict[str, Any]) -> bool:
    try:
        if manifest.get("schema") != "gradientharbor.integrity-baseline.v1":
            return False
        payload = _manifest_payload(str(manifest["filename"]), str(manifest["sha256"]), str(manifest["created_at"]))
        return verify_signature(payload, str(manifest["signature_ed25519"]))
    except (KeyError, TypeError, ValueError):
        return False


def compare_to_baseline(file_bytes: bytes, manifest: dict[str, Any]) -> dict[str, Any]:
    signature_ok = verify_signed_baseline(manifest)
    current = file_sha256(file_bytes)
    matches = signature_ok and current.lower() == str(manifest.get("sha256", "")).lower()
    return {"signature_valid": signature_ok, "hash_matches": current.lower() == str(manifest.get("sha256", "")).lower(),
            "match": matches, "current_sha256": current, "baseline_sha256": manifest.get("sha256")}


def sort_findings(findings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(findings, key=lambda item: (SEVERITY_ORDER.get(item.get("severity", "Info"), 9), item.get("category", ""), item.get("title", "")))
